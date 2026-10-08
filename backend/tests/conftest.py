from __future__ import annotations

import io
import uuid
from collections.abc import Sequence
from datetime import datetime, timezone

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from fpdf import FPDF
from langchain_core.embeddings import DeterministicFakeEmbedding
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage

from app.config import Settings
from app.container import Services, retrieval_config
from app.domain import ChunkDraft, Document, RetrievedChunk
from app.ingestion.service import ChunkingConfig, IngestionService
from app.main import create_app
from app.rag.chat import ChatService
from app.retrieval.reranker import NoopReranker
from app.retrieval.retriever import HybridRetriever

EMBED_DIM = 16


class FakeRepository:
    """In-memory stand-in for PostgresRepository. Search is naive but deterministic."""

    def __init__(self) -> None:
        self.documents: dict[uuid.UUID, Document] = {}
        self.hashes: dict[tuple[str, str], uuid.UUID] = {}
        self.chunks: list[tuple[uuid.UUID, ChunkDraft, Sequence[float]]] = []

    def ping(self) -> bool:
        return True

    def find_by_hash(self, collection: str, sha256: str) -> Document | None:
        doc_id = self.hashes.get((collection, sha256))
        return self.documents.get(doc_id) if doc_id else None

    def add_document(
        self,
        *,
        collection,
        filename,
        content_type,
        size_bytes,
        sha256,
        embedding_model,
        chunks,
        embeddings,
    ) -> Document:
        doc = Document(
            id=uuid.uuid4(),
            filename=filename,
            content_type=content_type,
            size_bytes=size_bytes,
            num_chunks=len(chunks),
            collection=collection,
            created_at=datetime.now(timezone.utc),
        )
        self.documents[doc.id] = doc
        self.hashes[(collection, sha256)] = doc.id
        for c, e in zip(chunks, embeddings, strict=True):
            self.chunks.append((doc.id, c, e))
        return doc

    def list_documents(self, collection: str) -> list[Document]:
        return [d for d in self.documents.values() if d.collection == collection]

    def get_document(self, document_id):
        return self.documents.get(document_id)

    def delete_document(self, document_id) -> bool:
        doc = self.documents.pop(document_id, None)
        if doc is None:
            return False
        self.chunks = [c for c in self.chunks if c[0] != document_id]
        return True

    def delete_collection(self, collection: str) -> int:
        ids = [d.id for d in self.list_documents(collection)]
        for i in ids:
            self.delete_document(i)
        return len(ids)

    def _retrieved(self, doc_id, chunk) -> RetrievedChunk:
        return RetrievedChunk(
            chunk_id=uuid.uuid5(doc_id, str(chunk.chunk_index)),
            document_id=doc_id,
            filename=self.documents[doc_id].filename,
            content=chunk.content,
            chunk_index=chunk.chunk_index,
            page=chunk.page,
        )

    def _in(self, collection):
        return [(d, c, e) for d, c, e in self.chunks if self.documents[d].collection == collection]

    def vector_search(self, collection, embedding, limit):
        def dot(e):
            return sum(a * b for a, b in zip(e, embedding, strict=True))

        ranked = sorted(self._in(collection), key=lambda t: dot(t[2]), reverse=True)[:limit]
        out = []
        for d, c, e in ranked:
            r = self._retrieved(d, c)
            r.vector_score = dot(e)
            out.append(r)
        return out

    def keyword_search(self, collection, query, limit):
        terms = {t.lower().strip("?.,") for t in query.split() if len(t) > 3}
        scored = []
        for d, c, _ in self._in(collection):
            score = sum(1 for t in terms if t in c.content.lower())
            if score:
                r = self._retrieved(d, c)
                r.keyword_score = float(score)
                scored.append(r)
        return sorted(scored, key=lambda r: r.keyword_score, reverse=True)[:limit]


def fake_llm(*answers: str) -> GenericFakeChatModel:
    return GenericFakeChatModel(messages=iter([AIMessage(content=a) for a in answers]))


def make_docx(*paragraphs: str) -> bytes:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(*pages: str) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.multi_cell(0, 6, text)
    return bytes(pdf.output())


@pytest.fixture
def settings() -> Settings:
    return Settings(
        _env_file=None,
        embedding_dim=EMBED_DIM,
        reranker="none",
        chunk_size=200,
        chunk_overlap=20,
        max_upload_mb=1,
        final_top_k=3,
    )


@pytest.fixture
def repository() -> FakeRepository:
    return FakeRepository()


@pytest.fixture
def embeddings() -> DeterministicFakeEmbedding:
    return DeterministicFakeEmbedding(size=EMBED_DIM)


@pytest.fixture
def ingestion(repository, embeddings, settings) -> IngestionService:
    return IngestionService(
        repository, embeddings, "fake-embed", ChunkingConfig(200, 20), settings.max_upload_bytes
    )


@pytest.fixture
def retriever(repository, embeddings, settings) -> HybridRetriever:
    return HybridRetriever(repository, embeddings, NoopReranker(), retrieval_config(settings))


@pytest.fixture
def make_client(settings, repository, ingestion, retriever):
    def _make(llm=None, errors=None) -> TestClient:
        def factory(s: Settings) -> Services:
            if errors:
                return Services(settings=s, errors=list(errors))
            return Services(
                settings=s,
                repository=repository,
                ingestion=ingestion,
                retriever=retriever,
                chat=ChatService(retriever, llm or fake_llm("Answer [1].")),
            )

        return TestClient(create_app(settings, services_factory=factory))

    return _make
