"""Text extraction for PDF and DOCX files."""

from __future__ import annotations

import io
import logging
import re
from pathlib import PurePath

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.domain import PageText

logger = logging.getLogger(__name__)

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_EXTENSIONS = {".pdf": PDF, ".docx": DOCX}
_MAGIC = {PDF: b"%PDF-", DOCX: b"PK\x03\x04"}


class IngestionError(ValueError):
    """Base class for user-facing ingestion problems (mapped to 4xx responses)."""


class UnsupportedFileError(IngestionError):
    pass


class EmptyDocumentError(IngestionError):
    pass


def detect_content_type(filename: str, data: bytes) -> str:
    """Return the MIME type from the extension, after checking the file signature."""
    ext = PurePath(filename).suffix.lower()
    content_type = _EXTENSIONS.get(ext)
    if content_type is None:
        raise UnsupportedFileError(
            f"Unsupported file type '{ext or filename}'. Upload PDF or DOCX."
        )
    if not data.startswith(_MAGIC[content_type]):
        raise UnsupportedFileError(f"'{filename}' does not look like a valid {ext} file")
    return content_type


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"-\n(?=\w)", "", text)  # re-join words hyphenated at line breaks
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(data: bytes) -> list[PageText]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnsupportedFileError("Password-protected PDFs are not supported")
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = normalize_text(page.extract_text() or "")
            if text:
                pages.append(PageText(text=text, page=number))
    except PdfReadError as exc:
        raise UnsupportedFileError(f"Could not read PDF: {exc}") from exc
    return pages


def parse_docx(data: bytes) -> list[PageText]:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several exception types for bad files
        raise UnsupportedFileError(f"Could not read DOCX: {exc}") from exc

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = normalize_text("\n".join(parts))
    return [PageText(text=text, page=None)] if text else []


def extract_text(filename: str, data: bytes) -> tuple[str, list[PageText]]:
    content_type = detect_content_type(filename, data)
    pages = parse_pdf(data) if content_type == PDF else parse_docx(data)
    if not pages:
        raise EmptyDocumentError(
            f"No extractable text found in '{filename}'. Scanned PDFs need OCR first."
        )
    return content_type, pages
