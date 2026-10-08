"""Generate the sample PDF/DOCX corpus used by the evaluation set.

All companies, products and figures are fictional. Run from backend/:

    python scripts/make_sample_docs.py

Requires the dev dependencies (fpdf2, python-docx).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from fpdf import FPDF

OUT_DIR = Path(__file__).resolve().parent.parent / "eval" / "sample_docs"

HANDBOOK = [
    ("h", "Brightpath Software - Employee Handbook (2026 edition)"),
    (
        "p",
        "This handbook summarises the policies that apply to all Brightpath Software employees. "
        "Brightpath Software is a fictional company used for demonstration purposes.",
    ),
    ("h", "1. Paid time off"),
    (
        "p",
        "Full-time employees accrue 1.75 days of paid time off (PTO) per month, which is 21 days "
        "per year. Part-time employees accrue PTO pro rata to their contracted hours.",
    ),
    (
        "p",
        "Up to 5 unused PTO days can be carried over into the next calendar year. Carried-over "
        "days expire on March 31 if they are not used.",
    ),
    (
        "p",
        "PTO requests longer than 5 consecutive working days must be submitted at least "
        "4 weeks in advance through the HR portal.",
    ),
    ("h", "2. Remote work"),
    (
        "p",
        "Employees may work remotely up to 3 days per week. Core collaboration hours are "
        "10:00 to 15:00 local time, during which employees are expected to be reachable.",
    ),
    (
        "p",
        "New remote-eligible employees receive a one-time home office stipend of $600 and a "
        "monthly internet allowance of $40.",
    ),
    ("h", "3. Expenses"),
    (
        "p",
        "Business expenses must be submitted within 30 days of purchase through the Ledgerly "
        "expense portal. Receipts are required for any expense over $25. Expenses above $500 "
        "require written approval from your manager before purchase.",
    ),
    ("h", "4. Parental leave"),
    (
        "p",
        "Primary caregivers receive 16 weeks of fully paid parental leave. Secondary caregivers "
        "receive 6 weeks of fully paid parental leave. Leave can be taken within 12 months "
        "of the birth or adoption.",
    ),
    ("h", "5. Learning and development"),
    (
        "p",
        "Every employee has an annual learning budget of $1,200 for courses, books and "
        "conferences. Unused learning budget does not roll over.",
    ),
    ("h", "6. Probation"),
    (
        "p",
        "New employees complete a 90-day probation period. During probation the notice period "
        "for either side is one week.",
    ),
    (
        "table",
        [
            ("Holiday", "Date"),
            ("New Year's Day", "January 1"),
            ("Founders' Day", "June 14"),
            ("Winter Break", "December 24 to December 31"),
        ],
    ),
]

MANUAL_PAGES = [
    [
        ("h", "Halcyon Aurora T2 Smart Thermostat - User Manual"),
        (
            "p",
            "Thank you for choosing the Aurora T2. This manual covers installation, Wi-Fi "
            "setup, troubleshooting and warranty. Halcyon Home and the Aurora T2 are fictional "
            "products used for demonstration purposes.",
        ),
        ("h", "In the box"),
        (
            "p",
            "Aurora T2 display unit, wall plate, trim plate, 4 screws and anchors, wire labels, "
            "and a C-wire adapter.",
        ),
    ],
    [
        ("h", "Installation"),
        (
            "p",
            "Before you begin, turn off power to your heating and cooling system at the circuit "
            "breaker. The Aurora T2 works with most 24V heating and cooling systems, including "
            "gas, oil, electric furnaces and heat pumps.",
        ),
        (
            "p",
            "The Aurora T2 is not compatible with line-voltage (120V or 240V) baseboard heaters. "
            "A C-wire (common wire) is required. If your system has no C-wire, install the "
            "included C-wire adapter at the furnace control board.",
        ),
        (
            "p",
            "Label each wire before disconnecting your old thermostat, then connect the wires to "
            "the matching terminals on the Aurora T2 wall plate.",
        ),
    ],
    [
        ("h", "Wi-Fi setup"),
        (
            "p",
            "The Aurora T2 connects to 2.4 GHz Wi-Fi networks only, using WPA2 or WPA3 security. "
            "5 GHz networks are not supported.",
        ),
        (
            "p",
            "To start pairing, press and hold the main button for 5 seconds until the LED blinks "
            "blue, then follow the steps in the Halcyon Home app.",
        ),
        ("h", "Error codes"),
        ("p", "E1: No power on the C-wire. Check the C-wire connection or install the adapter."),
        ("p", "E3: Temperature sensor fault. Restart the device; contact support if it persists."),
        ("p", "E7: Wi-Fi authentication failed. Re-enter your Wi-Fi password in the app."),
    ],
    [
        ("h", "Features"),
        (
            "p",
            "Eco mode keeps the temperature between 16 and 26 degrees Celsius while you are "
            "away. The built-in battery keeps the clock and schedule running for up to 8 hours "
            "during a power cut.",
        ),
        ("p", "To factory reset the thermostat, press and hold the main button for 15 seconds."),
        ("h", "Warranty"),
        (
            "p",
            "The Aurora T2 comes with a 2-year limited warranty from the date of purchase. "
            "The warranty does not cover damage caused by incorrect wiring.",
        ),
    ],
]

SECURITY_PAGES = [
    [
        ("h", "Brightpath Software - Information Security Policy"),
        (
            "p",
            "This policy applies to all employees, contractors and systems of Brightpath "
            "Software, a fictional company.",
        ),
        ("h", "Accounts and passwords"),
        (
            "p",
            "Passwords must be at least 14 characters long and must be stored in the company "
            "password manager. Multi-factor authentication (MFA) is mandatory for every "
            "account that supports it. Shared accounts are prohibited.",
        ),
        ("h", "Devices"),
        (
            "p",
            "All laptops must use full-disk encryption. Screens must lock automatically after "
            "5 minutes of inactivity. Lost or stolen devices must be reported immediately.",
        ),
    ],
    [
        ("h", "Incident response"),
        (
            "p",
            "Report any suspected security incident to security@brightpath.example within "
            "1 hour of discovery. The security team responds to Severity 1 incidents within "
            "15 minutes, around the clock.",
        ),
        ("h", "Data retention"),
        (
            "p",
            "Customer application logs are retained for 13 months and then deleted. Database "
            "backups are retained for 35 days.",
        ),
        ("h", "Access control and vendors"),
        (
            "p",
            "Access rights to production systems are reviewed quarterly by system owners. "
            "Vendors that process customer data must hold a current SOC 2 Type II report.",
        ),
    ],
]


def write_docx(path: Path) -> None:
    doc = Document()
    for kind, content in HANDBOOK:
        if kind == "h":
            doc.add_heading(content, level=1)
        elif kind == "p":
            doc.add_paragraph(content)
        else:
            table = doc.add_table(rows=0, cols=2)
            for left, right in content:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = left, right
    doc.save(path)


def write_pdf(path: Path, pages: list[list[tuple[str, str]]]) -> None:
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    for blocks in pages:
        pdf.add_page()
        for kind, text in blocks:
            if kind == "h":
                pdf.set_font("Helvetica", "B", 14)
                pdf.multi_cell(0, 8, text)
                pdf.ln(2)
            else:
                pdf.set_font("Helvetica", "", 11)
                pdf.multi_cell(0, 6, text)
                pdf.ln(3)
    pdf.output(str(path))


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_docx(OUT_DIR / "brightpath_employee_handbook.docx")
    write_pdf(OUT_DIR / "aurora_t2_user_manual.pdf", MANUAL_PAGES)
    write_pdf(OUT_DIR / "brightpath_security_policy.pdf", SECURITY_PAGES)
    for f in sorted(OUT_DIR.iterdir()):
        print(f"wrote {f.relative_to(OUT_DIR.parent.parent)} ({f.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
